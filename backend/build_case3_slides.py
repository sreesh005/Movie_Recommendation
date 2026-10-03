"""Builds the Case Study 3 slides from the class template (same layout as the GSW deck)."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts"
FIGS = ART / "report" / "figures"
SLIDE_FIGS = ART / "slides"
TEMPLATE = Path(r"C:\Users\kkr84\Downloads\Case Study Presentation Template (2).pptx")
OUT = Path(r"C:\Users\kkr84\Downloads\Case_Study_3_Later_Slides.pptx")

DARK = RGBColor(0x14, 0x11, 0x0E)
DARK_LINE = RGBColor(0x33, 0x2C, 0x24)
INK = RGBColor(0x1C, 0x19, 0x16)
BODY = RGBColor(0x4A, 0x45, 0x3F)
MUTED = RGBColor(0x8A, 0x82, 0x78)
CREAM = RGBColor(0xEF, 0xE6, 0xD6)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GOLD_TEXT = RGBColor(0xA9, 0x84, 0x12)
GOLD_TINT = RGBColor(0xFB, 0xF4, 0xDE)
PAGE = RGBColor(0xFA, 0xF8, 0xF5)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
CARD_ALT = RGBColor(0xF6, 0xF3, 0xEE)
LINE = RGBColor(0xE6, 0xE0, 0xD6)
GREEN = RGBColor(0x2E, 0x9E, 0x5B)
GREEN_TINT = RGBColor(0xEC, 0xF8, 0xF0)
RED = RGBColor(0xC2, 0x4D, 0x3A)
RED_TINT = RGBColor(0xFC, 0xEC, 0xE8)
BLUE_TINT = RGBColor(0xE6, 0xEE, 0xFB)
BLUE = RGBColor(0x2F, 0x5F, 0xB5)

HEAD = "Segoe UI Semibold"
TEXT = "Segoe UI"
TOTAL = 11


def load_metrics() -> dict:
    return json.loads((ART / "metrics.json").read_text(encoding="utf-8"))


def ndcg_chart(metrics: dict) -> Path:
    SLIDE_FIGS.mkdir(parents=True, exist_ok=True)
    wanted = [
        ("ItemKNN (original)", "ItemKNN (predicted stars)"),
        ("ItemKNN x log-count", "ItemKNN × log-count"),
        ("Popularity (count)", "Popularity (rating count)"),
        ("ImplicitMF (liked)", "ImplicitMF (liked ratings)"),
        ("Hybrid + ImplicitMF", "Hybrid + ImplicitMF (live)"),
    ]
    by_name = {row["algorithm"]: row for row in metrics["all"]}
    labels = [label for _, label in wanted]
    values = [by_name[key]["ndcg_at_20"] for key, _ in wanted]
    colors = ["#cfc6b8"] * (len(values) - 1) + ["#c9a227"]

    fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=200)
    bars = ax.barh(labels, values, color=colors, height=0.6)
    ax.invert_yaxis()
    for bar, v in zip(bars, values):
        ax.text(v + 0.004, bar.get_y() + bar.get_height() / 2, f"{v:.4f}", va="center", fontsize=9, color="#1c1916")
    ax.set_xlim(0, max(values) * 1.22)
    ax.set_xlabel("nDCG@20 (600 hidden ratings, 120 users)", fontsize=9, color="#4a453f")
    ax.tick_params(axis="both", labelsize=9, colors="#4a453f")
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#d8d0c4")
    ax.tick_params(axis="y", length=0)
    fig.tight_layout()
    out = SLIDE_FIGS / "ndcg_by_model.png"
    fig.savefig(out, facecolor="white")
    plt.close(fig)
    return out


def clear_slides(prs: Presentation) -> None:
    id_list = prs.slides._sldIdLst
    for sld_id in list(id_list):
        prs.part.drop_rel(sld_id.rId)
        id_list.remove(sld_id)


def blank(prs: Presentation, color: RGBColor):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color
    return slide


def box(slide, x, y, w, h, fill=None, line=None, radius=0.05, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def text(slide, x, y, w, h, runs, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, spacing=1.15):
    """Each paragraph is (text, size, color, font) or a list of those."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        parts = para if isinstance(para, list) else [para]
        for content, size, color, font in parts:
            r = p.add_run()
            r.text = content
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.name = font
    return tb


def header(slide, label: str, title: str, n: int) -> None:
    text(slide, 0.66, 0.42, 9, 0.3, [(label.upper(), 11, GOLD_TEXT, HEAD)])
    text(slide, 11.2, 0.42, 1.47, 0.3, [(f"{n:02d} / {TOTAL}", 11, INK, HEAD)], align=PP_ALIGN.RIGHT)
    text(slide, 0.66, 0.72, 12, 0.7, [(title, 28, INK, HEAD)])


def pill(slide, x, y, w, label, fill, color, h=0.32):
    shp = box(slide, x, y, w, h, fill=fill, radius=0.3)
    tf = shp.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(10)
    r.font.color.rgb = color
    r.font.name = HEAD
    return shp


def icon(slide, x, y, glyph, fill, color, size=0.42):
    shp = box(slide, x, y, size, size, fill=fill, radius=0.2)
    tf = shp.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = glyph
    r.font.size = Pt(14)
    r.font.color.rgb = color
    r.font.name = HEAD
    return shp


def check_list(slide, x, y, w, items, gap=0.62, size=13, marks=None):
    for i, item in enumerate(items):
        mark, fill = (marks[i] if marks else ("✓", GREEN))
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y + i * gap + 0.04), Inches(0.24), Inches(0.24))
        dot.fill.solid()
        dot.fill.fore_color.rgb = fill
        dot.line.fill.background()
        dot.shadow.inherit = False
        tf = dot.text_frame
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, 0)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = mark
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = CARD
        text(slide, x + 0.38, y + i * gap, w - 0.38, gap, [(item, size, INK, TEXT)])


def picture_fit(slide, path: Path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def notes(slide, body: str) -> None:
    slide.notes_slide.notes_text_frame.text = body


def build() -> Path:
    m = load_metrics()
    base, best = m["baseline_itemknn"], m["best"]
    rp = m["rating_prediction"]
    rmse_rows = {row["model"]: row for row in rp["rating_models"]}
    chart = ndcg_chart(m)

    prs = Presentation(str(TEMPLATE))
    clear_slides(prs)

    # 1. Title
    s = blank(prs, DARK)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.66), Inches(2.02), Inches(0.26), Inches(0.26))
    dot.fill.solid()
    dot.fill.fore_color.rgb = GOLD
    dot.line.fill.background()
    text(s, 1.04, 2.0, 8, 0.3, [("DAEN 400 CASE STUDY 3 PRESENTATION", 12, GOLD, HEAD)])
    text(s, 11.2, 2.0, 1.47, 0.3, [(f"01 / {TOTAL}", 11, MUTED, TEXT)], align=PP_ALIGN.RIGHT)
    text(s, 0.66, 2.45, 10.5, 1.6, [("Later: A Movie Recommender That Knows When to Watch", 40, CREAM, HEAD)], spacing=1.0)
    text(
        s, 0.66, 4.05, 9.5, 0.8,
        [("Building a custom LensKit recommender on MovieLens 32M with mood filters and watch-time suggestions", 16, MUTED, TEXT)],
    )
    line = s.shapes.add_connector(1, Inches(0.66), Inches(5.0), Inches(12.67), Inches(5.0))
    line.line.color.rgb = DARK_LINE
    meta = [("AUTHOR", "Sreesh Kanala", CREAM), ("DATASET", "MovieLens 32M", CREAM), ("FOCUS AREA", "Recommender Systems", GOLD)]
    for i, (label, value, color) in enumerate(meta):
        x = 0.66 + i * 2.6
        text(s, x, 5.22, 2.4, 0.25, [(label, 10, MUTED, HEAD)])
        text(s, x, 5.48, 2.4, 0.35, [(value, 14, color, HEAD)])
    box(s, 0, 7.42, 13.333, 0.08, fill=GOLD, shape=MSO_SHAPE.RECTANGLE)
    notes(s, "Case Study 3: building Later, a custom movie recommender on MovieLens 32M using LensKit.")

    # 2. Problem introduction
    s = blank(prs, PAGE)
    header(s, "Context & Objective", "Helping Users Decide What to Watch, and When", 2)
    box(s, 0.66, 2.35, 5.55, 2.55, fill=CARD_ALT, line=LINE)
    icon(s, 0.98, 2.67, "✕", RED_TINT, RED)
    text(s, 0.98, 3.25, 5.0, 0.4, [("Endless Scrolling (Status Quo)", 16, INK, HEAD)])
    text(
        s, 0.98, 3.7, 4.95, 1.1,
        [("MovieLens 32M has 87,585 movies. Generic popular lists ignore a user’s taste, mood, and how much time they have tonight.", 13, BODY, TEXT)],
    )
    text(s, 6.38, 3.35, 0.6, 0.5, [("➜", 22, GOLD_TEXT, HEAD)], align=PP_ALIGN.CENTER)
    box(s, 7.12, 2.35, 5.55, 2.55, fill=GREEN_TINT, line=RGBColor(0xCF, 0xEB, 0xD9))
    icon(s, 7.44, 2.67, "★", RGBColor(0xD3, 0xF0, 0xDE), GREEN)
    text(s, 7.44, 3.25, 5.0, 0.4, [("The Later Recommender", 16, INK, HEAD)])
    text(
        s, 7.44, 3.7, 4.95, 1.1,
        [("A LensKit recommender trained on 31.7 million ratings that builds a poster-style home page with mood rows and a suggested time to watch.", 13, BODY, TEXT)],
    )
    notes(s, "The problem is choice overload. Later narrows 87,585 movies to rows that fit the user's taste, mood, and time.")

    # 3. Analysis of the problem
    s = blank(prs, PAGE)
    header(s, "Problem Analysis", "Core Questions & Data Complications", 3)
    box(s, 0.66, 1.75, 12.0, 0.95, fill=GOLD_TINT, line=RGBColor(0xEE, 0xDF, 0xAE))
    text(
        s, 0.95, 1.9, 11.5, 0.7,
        [[
            ("Data Complication: ", 13, GOLD_TEXT, HEAD),
            ("MovieLens timestamps record when a movie was rated, not when it was watched, and about 73% of movies have fewer than 20 ratings, too few for the model to learn from.", 13, INK, TEXT),
        ]],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    questions = [
        ("Question 1: Recommendation Quality", "Can a LensKit model trained on the full dataset recommend movies a user actually likes?"),
        ("Question 2: Few Ratings", "How should the model handle a new user who has only rated a few movies in the survey?"),
        ("Question 3: Watch Time", "Can rating timestamps be used to suggest when a movie should be watched?"),
    ]
    for i, (q, body) in enumerate(questions):
        x = 0.66 + i * 4.07
        box(s, x, 3.05, 3.86, 2.7, fill=CARD, line=LINE)
        box(s, x, 3.05, 3.86, 0.08, fill=GOLD, shape=MSO_SHAPE.RECTANGLE)
        text(s, x + 0.3, 3.4, 3.3, 0.7, [(q, 15, INK, HEAD)])
        text(s, x + 0.3, 4.15, 3.3, 1.4, [(body, 13, BODY, TEXT)])
    notes(s, "Assumption: a rating means the user watched the movie. The timestamp limitation shapes how the watch-time feature is presented.")

    # 4. Approach
    s = blank(prs, PAGE)
    header(s, "Technical Framework", "Modeling Architecture & Hybrid Ranking", 4)
    box(s, 0.66, 1.7, 6.0, 5.2, fill=CARD, line=LINE)
    text(s, 0.98, 1.95, 5.4, 0.4, [("Approach", 16, INK, HEAD)])
    check_list(
        s, 0.98, 2.5, 5.45,
        [
            "Removed movies with fewer than 20 ratings: 23,350 movies, 200,763 users, 31.7M ratings.",
            "Trained ItemKNN first since it is built into LensKit and easy to explain.",
            "Ranking by ItemKNN’s predicted stars failed, so ImplicitMF was trained on liked ratings (4+ stars).",
            "Hybrid score = ImplicitMF + genre + popularity, weighted by how many movies the user has rated.",
            "Considered a two-tower network and graph walks; too heavy for one computer.",
        ],
        gap=0.86,
        size=12,
        marks=[("✓", GREEN), ("✓", GREEN), ("!", RGBColor(0xD9, 0x8A, 0x1C)), ("✓", GREEN), ("✕", MUTED)],
    )
    box(s, 6.9, 1.7, 5.77, 5.2, fill=CARD, line=LINE)
    text(s, 7.2, 1.95, 5.2, 0.4, [("Hybrid Weights by Movies Rated", 16, INK, HEAD)])
    picture_fit(s, FIGS / "fig2_hybrid_weights.png", 7.1, 2.45, 5.37, 4.25)
    notes(s, "After 5 or more ratings the weights are 0.62 ImplicitMF, 0.23 genre, 0.15 popularity. The survey asks for 8 ratings, so most users land there.")

    # 5. Findings: ranking
    s = blank(prs, PAGE)
    header(s, "Results — Ranking", f"Hybrid Model Finds {best['hits']} of 600 Hidden Ratings vs. {base['hits']}", 5)
    box(s, 0.66, 1.85, 4.9, 2.05, fill=CARD, line=LINE)
    text(s, 0.98, 2.1, 4.3, 0.3, [("ItemKNN (Predicted Stars) nDCG@20", 12, MUTED, HEAD)])
    text(s, 0.98, 2.45, 4.3, 0.8, [(f"{base['ndcg_at_20']:.4f}", 36, MUTED, HEAD)])
    text(s, 0.98, 3.3, 4.3, 0.4, [(f"{base['hits']} hits · recall@20 {base['recall_at_20']:.4f}", 12, MUTED, TEXT)])
    box(s, 0.66, 4.15, 4.9, 2.2, fill=CARD, line=GOLD)
    text(s, 0.98, 4.4, 4.3, 0.3, [("Hybrid + ImplicitMF nDCG@20", 12, GOLD_TEXT, HEAD)])
    text(s, 0.98, 4.75, 3.2, 0.8, [(f"{best['ndcg_at_20']:.4f}", 40, INK, HEAD)])
    lift = best["ndcg_at_20"] / base["ndcg_at_20"]
    pill(s, 3.95, 4.95, 1.35, f"{lift:.0f}× Lift", GOLD_TINT, GOLD_TEXT)
    text(s, 0.98, 5.7, 4.3, 0.4, [(f"{best['hits']} hits · recall@20 {best['recall_at_20']:.4f}", 12, BODY, TEXT)])
    box(s, 5.8, 1.85, 6.87, 4.5, fill=CARD, line=LINE)
    text(s, 6.1, 2.05, 6.3, 0.4, [("nDCG@20 by Model (same 120-user holdout)", 15, INK, HEAD)])
    picture_fit(s, chart, 6.0, 2.5, 6.47, 3.7)
    text(
        s, 0.66, 6.55, 12.0, 0.5,
        [(f"Random guessing would give recall@20 of {m['random_recall_at_20']:.5f}. ItemKNN kept picking lesser-known movies with high average ratings.", 12, BODY, TEXT)],
    )
    notes(s, "Five ratings were hidden from each of 120 users. A hit is a hidden movie that shows up in that user's top 20.")

    # 6. Findings: star ratings and likes
    s = blank(prs, PAGE)
    header(s, "Results — Ratings & Likes", "ItemKNN Still Predicts Stars Within About 2/3 of a Star", 6)
    box(s, 0.66, 1.85, 6.0, 4.6, fill=CARD, line=LINE)
    text(s, 0.98, 2.1, 5.4, 0.4, [("Star Prediction Error (RMSE, lower is better)", 15, INK, HEAD)])
    rows = [
        ("ItemKNN rating predictor", rmse_rows["ItemKNN rating predictor"]["rmse"], True),
        ("User mean", rmse_rows["User mean"]["rmse"], False),
        ("Item Bayesian average", rmse_rows["Item Bayesian average"]["rmse"], False),
        ("Global mean", rmse_rows["Global mean"]["rmse"], False),
    ]
    for i, (name, val, top) in enumerate(rows):
        y = 2.75 + i * 0.85
        box(s, 0.98, y, 5.36, 0.66, fill=GOLD_TINT if top else CARD_ALT, radius=0.15)
        text(s, 1.2, y, 3.6, 0.66, [(name, 13, INK, HEAD if top else TEXT)], anchor=MSO_ANCHOR.MIDDLE)
        label = f"{val:.3f} (Best)" if top else f"{val:.3f}"
        text(s, 4.2, y, 1.95, 0.66, [(label, 14, GOLD_TEXT if top else BODY, HEAD)], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    box(s, 6.9, 1.85, 5.77, 2.05, fill=CARD, line=LINE)
    text(s, 7.2, 2.1, 5.2, 0.3, [("Always Guess “Like” (Majority)", 12, MUTED, HEAD)])
    text(s, 7.2, 2.45, 5.2, 0.8, [(f"{rp['majority_class_accuracy'] * 100:.1f}%", 36, MUTED, HEAD)])
    text(s, 7.2, 3.3, 5.2, 0.4, [("Share of hidden ratings that were 4+ stars", 12, MUTED, TEXT)])
    box(s, 6.9, 4.15, 5.77, 2.3, fill=CARD, line=GOLD)
    text(s, 7.2, 4.4, 5.2, 0.3, [("Logistic Like Classifier Accuracy", 12, GOLD_TEXT, HEAD)])
    text(s, 7.2, 4.75, 3.2, 0.8, [(f"{rp['like_accuracy_logistic'] * 100:.1f}%", 40, INK, HEAD)])
    gain = (rp["like_accuracy_logistic"] - rp["majority_class_accuracy"]) * 100
    pill(s, 10.85, 4.95, 1.55, f"+{gain:.1f} pts", GOLD_TINT, GOLD_TEXT)
    text(
        s, 7.2, 5.7, 5.2, 0.6,
        [(f"ItemKNN thresholded at 4 stars: {rp['like_accuracy_itemknn'] * 100:.1f}%. MAE {rp['mae']:.3f} stars.", 12, BODY, TEXT)],
    )
    notes(s, "RMSE and like accuracy use the same 600 hidden ratings. These numbers judge star predictions, not the home page order.")

    # 7. Findings: interface
    s = blank(prs, PAGE)
    header(s, "Results — Final Product", "The Later Home Page", 7)
    box(s, 0.66, 1.7, 7.6, 5.2, fill=DARK, line=DARK_LINE)
    picture_fit(s, FIGS / "fig4_later_home.png", 0.8, 1.84, 7.32, 4.92)
    box(s, 8.5, 1.7, 4.17, 5.2, fill=CARD, line=LINE)
    text(s, 8.8, 1.95, 3.7, 0.4, [("What the User Sees", 16, INK, HEAD)])
    check_list(
        s, 8.8, 2.55, 3.65,
        [
            "Mood buttons with exclusive filters (comfort, date night, scary, laid-back, pay attention).",
            "“When to watch” label from genre watch windows.",
            "Half-star ratings and a “Not interested” button.",
            "Rows for tonight, the weekend, and “because you liked.”",
        ],
        gap=1.0,
        size=12,
    )
    notes(s, "Users rate eight movies in a survey, then get this home page. Ratings and Not interested update the rows.")

    # 8. Verification and validation
    s = blank(prs, PAGE)
    header(s, "Model Assurance", "Verification & Validation Framework", 8)
    box(s, 0.66, 1.7, 7.9, 5.2, fill=CARD, line=LINE)
    check_list(
        s, 0.98, 1.98, 7.3,
        [
            "Same 120-user, seed-42 holdout used for every model so results are comparable.",
            f"Checked against random (recall@20 {m['random_recall_at_20']:.5f}) and popularity baselines.",
            "RMSE compared against user-mean, item-average, and global-mean baselines.",
            "Every number pulled from artifacts/metrics.json written by the evaluation script.",
            f"Liked-only nDCG@20 of {best['liked_ndcg_at_20']:.4f} confirms the hits are movies users enjoyed.",
            "Mood rows made exclusive after the same movie showed up in conflicting rows.",
        ],
        gap=0.8,
        size=13,
    )
    box(s, 8.8, 1.7, 3.87, 2.45, fill=BLUE_TINT, line=RGBColor(0xC9, 0xD8, 0xF2))
    text(s, 9.1, 2.0, 3.3, 0.4, [("Verification", 18, BLUE, HEAD)])
    text(s, 9.1, 2.55, 3.3, 1.4, [("“Did we build the model right?”", 14, INK, TEXT), ("Same holdout, baselines, and metrics checks.", 12, MUTED, TEXT)])
    box(s, 8.8, 4.45, 3.87, 2.45, fill=GOLD_TINT, line=RGBColor(0xEE, 0xDF, 0xAE))
    text(s, 9.1, 4.75, 3.3, 0.4, [("Validation", 18, GOLD_TEXT, HEAD)])
    text(s, 9.1, 5.3, 3.3, 1.4, [("“Did we build the right model?”", 14, INK, TEXT), ("Hits are liked movies, and rows match the user’s mood.", 12, MUTED, TEXT)])
    notes(s, "Verification checks the code and metrics. Validation checks that the recommender solves the real problem of suggesting movies users like.")

    # 9. Conclusions
    s = blank(prs, PAGE)
    header(s, "Executive Takeaways", "Conclusions & Root Cause Analysis", 9)
    concl = [
        (
            "CONCLUSION 1", BLUE_TINT, BLUE, "The Ranking Method Mattered Most",
            f"ItemKNN predicts stars well (RMSE {m['rmse']:.3f}), but sorting by predicted stars gave an nDCG@20 of "
            f"{base['ndcg_at_20']:.4f}. Learning from the movies users liked and adding popularity raised it to "
            f"{best['ndcg_at_20']:.4f}.",
        ),
        (
            "CONCLUSION 2", GOLD_TINT, GOLD_TEXT, "Watch Times Reflect Rating Habits",
            "Rating peaks for almost every genre landed on Sunday around 3 p.m. The watch times are a useful product "
            "feature, but they show when people rate movies, not when they watch them.",
        ),
    ]
    for i, (tag, fill, color, title, body) in enumerate(concl):
        x = 0.66 + i * 6.12
        box(s, x, 1.7, 5.88, 5.2, fill=CARD_ALT, line=LINE)
        pill(s, x + 0.35, 3.0, 1.55, tag, fill, color)
        text(s, x + 0.35, 3.5, 5.2, 0.5, [(title, 17, INK, HEAD)])
        text(s, x + 0.35, 4.1, 5.15, 1.6, [(body, 13, BODY, TEXT)])
    notes(s, "Main takeaway: what the model is asked to rank matters more than which algorithm is used.")

    # 10. Recommendations
    s = blank(prs, PAGE)
    header(s, "Action Plan", "Recommended Roadmap for Later", 10)
    recs = [
        ("Adopt the Hybrid and Test at Scale", "Use the hybrid model as the main recommender. Repeat the test on 1,000 to 2,000 users and compare against SVD with a paired significance test."),
        ("Test Mood Buttons with Real Users", "Turn each mood filter and the “Not interested” penalty on and off in the test, and run a one-week study with 20 to 30 students."),
        ("Collect Viewing Data", "Log play, finish, and quit times to replace rating timestamps, and feed finished movies back into ImplicitMF as training data."),
    ]
    for i, (title, body) in enumerate(recs):
        x = 0.66 + i * 4.07
        box(s, x, 1.7, 3.86, 4.6, fill=CARD, line=LINE)
        icon(s, x + 0.3, 2.0, str(i + 1), GOLD_TINT, GOLD_TEXT, size=0.5)
        text(s, x + 0.3, 2.75, 3.3, 0.8, [(title, 16, INK, HEAD)])
        text(s, x + 0.3, 3.65, 3.3, 2.4, [(body, 13, BODY, TEXT)])
    text(
        s, 0.66, 6.5, 12.0, 0.4,
        [(f"Star predictions (RMSE {m['rmse']:.3f}) and like labels ({rp['like_accuracy_logistic'] * 100:.1f}%) can go on movie cards, but should not decide the home page order.", 12, BODY, TEXT)],
    )
    notes(s, "Next steps: a bigger test, measure the mood features, and collect real viewing data.")

    # 11. Questions
    s = blank(prs, DARK)
    text(s, 11.2, 0.42, 1.47, 0.3, [(f"{TOTAL} / {TOTAL}", 11, MUTED, TEXT)], align=PP_ALIGN.RIGHT)
    text(s, 0.66, 2.9, 12.0, 1.2, [("Questions", 54, CREAM, HEAD)], align=PP_ALIGN.CENTER)
    text(s, 0.66, 4.1, 12.0, 0.5, [("Later · MovieLens 32M · LensKit", 16, GOLD, TEXT)], align=PP_ALIGN.CENTER)
    box(s, 0, 7.42, 13.333, 0.08, fill=GOLD, shape=MSO_SHAPE.RECTANGLE)

    try:
        prs.save(str(OUT))
        return OUT
    except PermissionError:
        alt = OUT.with_name(OUT.stem + "_v2.pptx")
        prs.save(str(alt))
        return alt


if __name__ == "__main__":
    print("Saved", build())
